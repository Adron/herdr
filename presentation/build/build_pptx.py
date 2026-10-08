"""Build presentation/herding-agents.pptx from slides.json.

Uses python-pptx's stock "Office Theme" (master + layouts + placeholders), widened
to 16:9, so the deck stays a plain default theme that PowerPoint's Design tab can
swap in one click. Speaker notes come straight from slides.json.

    python3 -m pip install python-pptx pillow
    python3 presentation/build/build_pptx.py
"""
import copy
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Inches

ROOT = Path(__file__).resolve().parents[2]
SPEC = json.loads((Path(__file__).parent / "slides.json").read_text())
IMAGES = ROOT / "images"
OUT = ROOT / "presentation" / "herding-agents.pptx"

prs = Presentation()  # stock Office Theme, 4:3

# Widen to 16:9 (13.333" x 7.5") and stretch master/layout placeholders to match.
old_w = prs.slide_width
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
scale = prs.slide_width / old_w
# layouts first: their inherited values must be read before the master changes
for shapes in [l.shapes for l in prs.slide_layouts] + [prs.slide_master.shapes]:
    for shp in shapes:
        if shp.left is not None:
            # read all four first: assigning one creates an xfrm that zeroes the rest
            x, y, w, h = shp.left, shp.top, shp.width, shp.height
            shp.left, shp.top = Emu(int(x * scale)), Emu(y)
            shp.width, shp.height = Emu(int(w * scale)), Emu(h)

L = {l.name: l for l in prs.slide_layouts}
W, H = prs.slide_width, prs.slide_height
MARGIN = Inches(0.5)


def drop_empty_placeholders(slide):
    for ph in list(slide.placeholders):
        if not ph.has_text_frame or not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def fit_picture(slide, path, left, top, max_w, max_h):
    iw, ih = Image.open(path).size
    ratio = min(max_w / iw, max_h / ih)
    w, h = int(iw * ratio), int(ih * ratio)
    x = left + (max_w - w) // 2
    y = top + (max_h - h) // 2
    slide.shapes.add_picture(str(path), x, y, w, h)


def add_author_date(slide):
    """Copy the layout's footer placeholder onto the slide and fill in author and date."""
    footer = next(p for p in slide.slide_layout.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    slide.shapes._spTree.append(copy.deepcopy(footer._element))
    ph = next(p for p in slide.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    ph.text_frame.text = SPEC["meta"]["author_date"]


def set_lines(text_frame, text):
    lines = text.split("\n")
    text_frame.text = lines[0]
    for line in lines[1:]:
        text_frame.add_paragraph().text = line


for s in SPEC["slides"]:
    t = s["type"]
    if t == "title":
        slide = prs.slides.add_slide(L["Title Slide"])
        slide.shapes.title.text = s["title"]
        set_lines(slide.placeholders[1].text_frame, s.get("subtitle", ""))
        add_author_date(slide)
    elif t in ("statement", "bigfact"):
        slide = prs.slides.add_slide(L["Section Header"])
        set_lines(slide.shapes.title.text_frame, s["title"])
        if s.get("subtitle"):
            slide.placeholders[1].text_frame.text = s["subtitle"]
    elif t == "bullets":
        slide = prs.slides.add_slide(L["Title and Content"])
        slide.shapes.title.text = s["title"]
        set_lines(slide.placeholders[1].text_frame, "\n".join(s["bullets"]))
    elif t in ("image", "image-side"):
        slide = prs.slides.add_slide(L["Title Only"])
        title = slide.shapes.title
        title.text = s["title"]
        img = IMAGES / s["image"]
        if t == "image-side":
            # tall, narrow screenshot: title on the left, image on the right
            title.left, title.top = MARGIN, Inches(2.5)
            title.width, title.height = Inches(7.5), Inches(2.5)
            fit_picture(slide, img, Inches(8.5), MARGIN, W - Inches(9.0), H - 2 * MARGIN)
        else:
            top = title.top + title.height + Inches(0.15)
            fit_picture(slide, img, MARGIN, top, W - 2 * MARGIN, H - top - MARGIN)
    else:
        raise ValueError(f"unknown slide type {t}")

    drop_empty_placeholders(slide)
    slide.notes_slide.notes_text_frame.text = s["notes"]

prs.core_properties.title = SPEC["meta"]["title"]
prs.core_properties.author = SPEC["meta"]["author"]
prs.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} ({len(SPEC['slides'])} slides)")
